import json
from pathlib import Path
# from uuid import uuid4
import hashlib

import pandas as pd
from docx import Document as DocxDocument
from pypdf import PdfReader

from app.rag.ingestion.models import Document


SUPPORTED_EXTENSIONS = {
    ".txt",
    ".md",
    ".pdf",
    ".docx",
    ".csv",
    ".json",
    ".xlsx",
}


def load_text_file(path: Path) -> str:
    return path.read_text(
        encoding="utf-8",
        errors="ignore",
    )


def load_pdf(path: Path) -> str:
    reader = PdfReader(str(path))

    pages = []

    for page in reader.pages:
        text = page.extract_text() or ""

        if text.strip():
            pages.append(text)

    return "\n".join(pages)


def load_docx(path: Path) -> str:
    document = DocxDocument(str(path))

    paragraphs = [
        paragraph.text
        for paragraph in document.paragraphs
        if paragraph.text.strip()
    ]

    return "\n".join(paragraphs)


def load_csv(path: Path) -> str:
    dataframe = pd.read_csv(path)

    return dataframe.to_csv(
        index=False
    )


def load_json(path: Path) -> str:
    with path.open(
        "r",
        encoding="utf-8",
    ) as file:
        data = json.load(file)

    return json.dumps(
        data,
        ensure_ascii=False,
        indent=2,
    )


def generate_document_id(path: Path) -> str:
    """
    Generate a stable document ID from the source path.
    The same document receives the same ID on every indexing run.
    """
    normalized_path = str(path.resolve()).lower()

    return hashlib.sha256(
        normalized_path.encode("utf-8")
    ).hexdigest()[:24]

def load_xlsx(path: Path) -> str:
    workbook = pd.ExcelFile(path)

    sections = []

    for sheet_name in workbook.sheet_names:

        dataframe = pd.read_excel(
            path,
            sheet_name=sheet_name,
        )

        sections.append(
            f"Sheet: {sheet_name}\n"
            f"{dataframe.to_csv(index=False)}"
        )

    return "\n\n".join(sections)


def load_document(path: Path) -> Document:

    extension = path.suffix.lower()

    if extension in {".txt", ".md"}:
        content = load_text_file(path)

    elif extension == ".pdf":
        content = load_pdf(path)

    elif extension == ".docx":
        content = load_docx(path)

    elif extension == ".csv":
        content = load_csv(path)

    elif extension == ".json":
        content = load_json(path)

    elif extension == ".xlsx":
        content = load_xlsx(path)

    else:
        raise ValueError(
            f"Unsupported document type: {extension}"
        )

    return Document(
        document_id=generate_document_id(path),
        source=str(path),
        title=path.stem,
        content=content.strip(),
        metadata={
            "filename": path.name,
            "extension": extension,
        },
    )